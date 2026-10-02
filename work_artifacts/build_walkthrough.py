from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = ROOT / "work_artifacts" / "walkthrough_assets"
OUTPUT = ROOT / "deliverables" / "Genedrift_Editorial_Platform_Client_Walkthrough.docx"

PURPLE = "#6F9187"
DARK = "#2C3B38"
INK = "#293330"
MUTED = "#6F7B77"
LAVENDER = "#EAF1EE"
PALE = "#F7F8F5"
BORDER = "#D8E1DE"
WHITE = "#FFFFFF"
GREEN = "#58796F"
AMBER = "#A58268"
RED = "#9A645B"
MIST_BLUE = "#EAF1F4"
WARM_SAND = "#F4EFE7"
SOFT_CLAY = "#C49F87"

# Base preset: compact_reference_guide.
# Named override: soft_editorial — Arial body, Georgia headings, sage/mist palette,
# quiet table headers, borderless running furniture and an editorial cover.
FONT = "Arial"
HEADING_FONT = "Georgia"
DIAGRAM_FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"
DIAGRAM_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value.lstrip("#"))


def set_run_font(run, size=11, color=INK, bold=False, italic=False, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.lstrip("#"))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, spec in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), spec.get("val", "single"))
        edge.set(qn("w:sz"), str(spec.get("sz", 6)))
        edge.set(qn("w:color"), spec.get("color", BORDER).lstrip("#"))


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_repeat_table_header(row):
    repeat_header(row)


def shade_paragraph(paragraph, fill=LAVENDER, border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill.lstrip("#"))
    p_pr.append(shd)
    if border:
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border.lstrip("#"))
        p_bdr.append(left)
        p_pr.append(p_bdr)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, size=8.5, color=MUTED)


def set_image_alt(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def paragraph(doc, text="", size=11, color=INK, bold=False, italic=False, after=6,
              before=0, line=1.25, align=WD_ALIGN_PARAGRAPH.LEFT, keep=False, name=FONT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic, name=name)
    return p


def add_kicker(doc, text):
    p = paragraph(doc, text.upper(), size=8.2, color=PURPLE, bold=True, after=7, line=1.0, keep=True)
    p.runs[0].font.all_caps = True
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = False
    run = p.add_run(text)
    return p


def add_page_title(doc, number, title, lead=None):
    kicker = add_kicker(doc, f"{number:02d} / Client field guide")
    kicker.paragraph_format.page_break_before = False
    add_heading(doc, title, 1)
    if lead:
        paragraph(doc, lead, size=11.8, color=MUTED, after=18, line=1.35)


def add_callout(doc, label, text, fill=PALE, accent=PURPLE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(11)
    p.paragraph_format.left_indent = Inches(0.14)
    p.paragraph_format.right_indent = Inches(0.14)
    p.paragraph_format.line_spacing = 1.25
    shade_paragraph(p, fill, accent)
    r1 = p.add_run(f"{label.upper()}  ")
    set_run_font(r1, size=8.2, color=accent, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.2, color=INK, bold=False)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "50")
    num_pr.extend([ilvl, num_id])
    p_pr.append(num_pr)
    set_run_font(p.add_run(text), size=10.2, color=INK)
    return p


def add_numbered(doc, text, num_id=51):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_element = OxmlElement("w:numId")
    num_element.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_element])
    p_pr.append(num_pr)
    set_run_font(p.add_run(text), size=10.6, color=INK)
    return p


def add_figure(doc, path, caption, alt, width=6.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    set_image_alt(shape, caption, alt)
    cap = paragraph(doc, caption, size=8.5, color=MUTED, italic=True, after=12, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
    cap.paragraph_format.keep_with_next = False


def add_table(doc, headers, rows, widths, header_fill=LAVENDER, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths, 120)
    repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        set_run_font(p.add_run(header), size=8.3, color=DARK, bold=True)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        set_table_geometry(table, widths, 120)
        for index, value in enumerate(values):
            cell = cells[index]
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            set_run_font(p.add_run(str(value)), size=font_size, color=INK, bold=(index == 0))
            borders = {e: {"color": BORDER, "sz": 3} for e in ("top", "left", "bottom", "right")}
            set_cell_border(cell, **borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def populate_running_header_footer(section):
    for header in (section.header, section.even_page_header, section.first_page_header):
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        for child in list(hp._p):
            hp._p.remove(child)
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_after = Pt(0)
        set_run_font(hp.add_run("GENEDRIFT   /   EDITORIAL PLATFORM"), size=7.5, color=MUTED, bold=False)

    for footer in (section.footer, section.even_page_footer, section.first_page_footer):
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        for child in list(fp._p):
            fp._p.remove(child)
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.paragraph_format.space_before = Pt(0)
        set_run_font(fp.add_run("CLIENT WALKTHROUGH   /   SEPTEMBER 2026   /   "), size=7.4, color=MUTED)
        add_page_field(fp)


def page_break(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = False
    populate_running_header_footer(section)


def diagram_font(size, bold=False):
    return ImageFont.truetype(DIAGRAM_BOLD if bold else DIAGRAM_FONT, size)


def rounded(draw, xy, fill, outline=None, radius=24, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def centered_text(draw, xy, text, font, fill, max_width=None, spacing=6):
    x1, y1, x2, y2 = xy
    lines = [text]
    if max_width:
        estimated = max(12, int(max_width / (font.size * 0.52)))
        lines = wrap(text, estimated) or [text]
    total_h = sum(draw.textbbox((0, 0), line, font=font)[3] for line in lines) + spacing * (len(lines) - 1)
    y = y1 + (y2 - y1 - total_h) / 2
    for line in lines:
        box = draw.textbbox((0, 0), line, font=font)
        draw.text((x1 + (x2 - x1 - (box[2] - box[0])) / 2, y), line, font=font, fill=fill)
        y += box[3] - box[1] + spacing


def arrow(draw, start, end, color=PURPLE, width=7):
    draw.line([start, end], fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - 18 * direction, y2 - 12), (x2 - 18 * direction, y2 + 12)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 12, y2 - 18 * direction), (x2 + 12, y2 - 18 * direction)]
    draw.polygon(points, fill=color)


def create_architecture_diagram(path):
    image = Image.new("RGB", (1800, 820), WHITE)
    draw = ImageDraw.Draw(image)
    title = diagram_font(34, True)
    body = diagram_font(24)
    small = diagram_font(19)
    boxes = [
        ((80, 190, 400, 650), WARM_SAND, DARK, "Editorial team", ["Authors", "Reviewers", "Publishers", "Editorial Admin"]),
        ((500, 190, 850, 650), LAVENDER, DARK, "Zoho Creator", ["Dashboard", "Article workspace", "Review + audit", "Publishing controls"]),
        ((950, 190, 1300, 650), MIST_BLUE, DARK, "Zoho Catalyst", ["Validates authority", "Stores immutable content", "Processes media", "Returns signed result"]),
        ((1400, 190, 1720, 650), PALE, DARK, "Public website", ["Insights listing", "Article pages", "Search + filters", "SEO delivery"]),
    ]
    for xy, fill, text_color, heading, items in boxes:
        rounded(draw, xy, fill, outline=BORDER, radius=20, width=2)
        draw.text((xy[0] + 30, xy[1] + 34), heading, font=title, fill=text_color)
        y = xy[1] + 115
        for item in items:
            draw.ellipse((xy[0] + 32, y + 7, xy[0] + 44, y + 19), fill=PURPLE)
            draw.text((xy[0] + 62, y), item, font=body, fill=text_color)
            y += 72
    arrow(draw, (410, 420), (482, 420), color=GREEN, width=5)
    arrow(draw, (860, 420), (932, 420), color=GREEN, width=5)
    arrow(draw, (1310, 420), (1382, 420), color=GREEN, width=5)
    draw.text((80, 85), "One governed path from editorial work to public delivery", font=diagram_font(42, True), fill=DARK)
    draw.text((80, 710), "Creator controls the work. Catalyst controls publication integrity. The website serves the approved result.", font=small, fill=MUTED)
    image.save(path, quality=95)


def create_dashboard_diagram(path):
    image = Image.new("RGB", (1800, 980), PALE)
    draw = ImageDraw.Draw(image)
    rounded(draw, (65, 60, 1735, 920), WHITE, BORDER, 22, 2)
    draw.text((115, 95), "GENEDRIFT INSIGHTS", font=diagram_font(34, True), fill=DARK)
    draw.text((115, 145), "Role-aware editorial dashboard", font=diagram_font(22), fill=MUTED)
    rounded(draw, (112, 220, 560, 290), PALE, BORDER, 16, 2)
    rounded(draw, (120, 228, 240, 282), LAVENDER, None, 12)
    labels = [("MINE", 147, PURPLE), ("QUEUE", 285, MUTED), ("ALL", 438, MUTED)]
    for text_value, x, color in labels:
        draw.text((x, 241), text_value, font=diagram_font(19, True), fill=color)
    rounded(draw, (610, 220, 1420, 290), WHITE, BORDER, 16, 2)
    draw.text((650, 241), "Search titles, people, categories or decisions", font=diagram_font(19), fill=MUTED)
    rounded(draw, (1460, 220, 1688, 290), PURPLE, None, 16)
    centered_text(draw, (1460, 220, 1688, 290), "REFRESH", diagram_font(18, True), WHITE)
    filters = [(112, 325, 420, 388, "ALL WORK"), (445, 325, 735, 388, "STATE"), (760, 325, 1075, 388, "CATEGORY"), (1100, 325, 1415, 388, "SORT"), (1440, 325, 1688, 388, "CLEAR")]
    for x1, y1, x2, y2, label in filters:
        rounded(draw, (x1, y1, x2, y2), WHITE, BORDER, 12, 2)
        centered_text(draw, (x1, y1, x2, y2), label, diagram_font(16, True), DARK)
    metrics = [("Review inbox", "4"), ("Active articles", "7"), ("Review history", "12"), ("Ready next", "3")]
    x = 112
    metric_fills = [LAVENDER, MIST_BLUE, WARM_SAND, PALE]
    for index, (label, value) in enumerate(metrics):
        rounded(draw, (x, 430, x + 365, 555), metric_fills[index], BORDER, 16, 2)
        draw.text((x + 24, 455), value, font=diagram_font(35, True), fill=PURPLE)
        draw.text((x + 95, 463), label, font=diagram_font(19, True), fill=DARK)
        x += 392
    panels = [(112, 605, 860, 845, "REVIEW INBOX", "Assigned and claimable review work"), (900, 605, 1688, 845, "ARTICLE / PUBLISHING WORK", "Drafts, approvals, schedules and public status")]
    for x1, y1, x2, y2, title, subtitle in panels:
        rounded(draw, (x1, y1, x2, y2), WHITE, BORDER, 16, 2)
        draw.text((x1 + 28, y1 + 28), title, font=diagram_font(21, True), fill=DARK)
        draw.text((x1 + 28, y1 + 72), subtitle, font=diagram_font(18), fill=MUTED)
        for i in range(2):
            y = y1 + 125 + i * 54
            draw.line((x1 + 28, y, x2 - 28, y), fill=BORDER, width=2)
    image.save(path, quality=95)


def create_review_diagram(path):
    image = Image.new("RGB", (1800, 960), WHITE)
    draw = ImageDraw.Draw(image)
    draw.text((75, 70), "Two review models, one controlled publishing path", font=diagram_font(42, True), fill=DARK)
    columns = [
        (75, 175, 865, 855, "STANDARD REVIEW", "One distinct approval", [("Draft", WARM_SAND), ("Submit", MIST_BLUE), ("Reviewer approves", LAVENDER), ("Approved", GREEN)]),
        (935, 175, 1725, 855, "REGULATED REVIEW", "Two distinct approvals", [("Draft", WARM_SAND), ("Submit", MIST_BLUE), ("Reviewer A approves", LAVENDER), ("Still In Review", PALE), ("Reviewer B approves", LAVENDER), ("Approved", GREEN)]),
    ]
    for x1, y1, x2, y2, heading, subtitle, steps in columns:
        rounded(draw, (x1, y1, x2, y2), PALE, BORDER, 22, 2)
        draw.text((x1 + 38, y1 + 34), heading, font=diagram_font(27, True), fill=PURPLE)
        draw.text((x1 + 38, y1 + 80), subtitle, font=diagram_font(21), fill=MUTED)
        top = y1 + 145
        height = 70 if len(steps) > 4 else 92
        gap = 22
        for index, (label, fill) in enumerate(steps):
            box = (x1 + 105, top, x2 - 105, top + height)
            rounded(draw, box, fill, BORDER if fill != GREEN else GREEN, 16, 2)
            centered_text(draw, box, label, diagram_font(21, True), WHITE if fill == GREEN else DARK, max_width=box[2]-box[0]-40)
            if index < len(steps) - 1:
                arrow(draw, ((x1 + x2) // 2, top + height + 2), ((x1 + x2) // 2, top + height + gap - 4), GREEN, 4)
            top += height + gap
    draw.text((75, 895), "At any review stage, Request Changes creates a new draft revision while preserving the submitted revision and its feedback.", font=diagram_font(20), fill=MUTED)
    image.save(path, quality=95)


def create_publication_diagram(path):
    image = Image.new("RGB", (1800, 790), PALE)
    draw = ImageDraw.Draw(image)
    draw.text((75, 62), "Publication is a separate, verified step", font=diagram_font(42, True), fill=DARK)
    rounded(draw, (80, 185, 390, 335), WARM_SAND, BORDER, 18, 2)
    centered_text(draw, (80, 185, 390, 335), "Approved\nrevision", diagram_font(25, True), DARK, spacing=8)
    rounded(draw, (80, 440, 390, 590), LAVENDER, BORDER, 18, 2)
    centered_text(draw, (80, 440, 390, 590), "Scheduled\nrevision", diagram_font(25, True), DARK, spacing=8)
    rounded(draw, (535, 300, 850, 475), LAVENDER, BORDER, 18, 2)
    centered_text(draw, (535, 300, 850, 475), "Creator job\n+ signed handoff", diagram_font(24, True), DARK, spacing=8)
    rounded(draw, (1000, 300, 1325, 475), MIST_BLUE, BORDER, 18, 2)
    centered_text(draw, (1000, 300, 1325, 475), "Catalyst validates\n+ stores immutable files", diagram_font(23, True), DARK, spacing=8)
    rounded(draw, (1470, 300, 1730, 475), WHITE, BORDER, 18, 2)
    centered_text(draw, (1470, 300, 1730, 475), "Public\nInsights", diagram_font(25, True), DARK, spacing=8)
    arrow(draw, (400, 260), (525, 360), GREEN, 5)
    arrow(draw, (400, 515), (525, 420), GREEN, 5)
    arrow(draw, (860, 388), (990, 388), GREEN, 5)
    arrow(draw, (1335, 388), (1460, 388), GREEN, 5)
    draw.text((102, 620), "PUBLISH NOW", font=diagram_font(18, True), fill=PURPLE)
    draw.text((102, 652), "Immediate processing", font=diagram_font(18), fill=MUTED)
    draw.text((1000, 540), "Callback confirms one truthful final state in Creator", font=diagram_font(18, True), fill=DARK)
    draw.text((1000, 580), "Retry and Reconcile reuse the same idempotent job", font=diagram_font(18), fill=MUTED)
    image.save(path, quality=95)


def install_numbering(doc):
    numbering = doc.part.numbering_part.element
    for abstract_id, num_id, fmt, text_value in ((50, 50, "bullet", "•"), (51, 51, "decimal", "%1.")):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text_value)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "270")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)
    for num_id in (52, 53):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), "51")
        num.append(abstract_ref)
        lvl_override = OxmlElement("w:lvlOverride")
        lvl_override.set(qn("w:ilvl"), "0")
        start_override = OxmlElement("w:startOverride")
        start_override.set(qn("w:val"), "1")
        lvl_override.append(start_override)
        num.append(lvl_override)
        numbering.append(num)


def configure_document(doc):
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        1: (23, DARK, 18, 10),
        2: (15, PURPLE, 14, 7),
        3: (12, DARK, 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        font_name = HEADING_FONT if level in (1, 2) else FONT
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn("w:ascii"), font_name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.05 if level == 1 else 1.15
        style.paragraph_format.keep_with_next = True

    populate_running_header_footer(section)
    install_numbering(doc)


def cover_page(doc, architecture_path):
    paragraph(doc, "Genedrift Editorial Publishing Platform", size=8.6, color=PURPLE, bold=True, after=82, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(
        doc,
        "A clear path from\nidea to publication",
        size=32,
        color=DARK,
        bold=False,
        after=12,
        line=1.05,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        name=HEADING_FONT,
    )
    paragraph(
        doc,
        "A calm, practical guide to the Creator workspace, review decisions, publishing controls and public delivery.",
        size=12.4,
        color=MUTED,
        after=26,
        line=1.35,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_figure(
        doc,
        architecture_path,
        "Platform orientation — editorial control, trusted publication and public delivery.",
        "Four-stage platform map showing the editorial team, Zoho Creator, Zoho Catalyst and the public website.",
        width=5.75,
    )
    add_callout(doc, "Designed for clarity", "Each person sees the work relevant to their role, while every important decision remains traceable in the audit history.", fill=PALE)
    paragraph(doc, "Prepared for Genedrift stakeholders  ·  Version 1.0  ·  September 2026", size=8.8, color=MUTED, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)


def build_document():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    architecture = ASSET_DIR / "platform_architecture.png"
    dashboard = ASSET_DIR / "dashboard_orientation.png"
    review = ASSET_DIR / "review_models.png"
    publication = ASSET_DIR / "publication_path.png"
    create_architecture_diagram(architecture)
    create_dashboard_diagram(dashboard)
    create_review_diagram(review)
    create_publication_diagram(publication)

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Genedrift Editorial Publishing Platform — Client Walkthrough"
    doc.core_properties.subject = "Client guide to the Genedrift Creator editorial and publishing workflow"
    doc.core_properties.author = "Genedrift"
    doc.core_properties.keywords = "Genedrift, editorial workflow, Zoho Creator, publishing, review"

    cover_page(doc, architecture)

    add_page_title(doc, 1, "The platform in one view", "Genedrift combines a role-aware editorial workspace with a controlled publishing engine and a modern public Insights experience.")
    add_figure(doc, architecture, "The end-to-end platform map.", "Editorial roles work in Creator; Catalyst validates and stores immutable publications; the public website serves the approved result.", width=5.75)
    add_heading(doc, "How the pieces work together", 2)
    add_bullet(doc, "Zoho Creator manages content, roles, reviews, status and audit history.")
    add_bullet(doc, "The Article Workspace supports day-to-day writing, reviewing and publishing.")
    add_bullet(doc, "Zoho Catalyst validates the approved revision, stores immutable files and returns a signed result.")
    add_bullet(doc, "The public website displays only content that completed the governed path.")
    add_callout(doc, "Core principle", "Approval confirms that content is ready. Publication is a separate Publisher decision.")
    page_break(doc)

    add_page_title(doc, 2, "Who does what", "Access is role-based. Combined roles are possible, but each action is still checked against its own permission.")
    add_table(
        doc,
        ["Role", "Primary responsibilities", "Key guardrail"],
        [
            ["Editorial Author", "Create drafts, write, add media and metadata, submit, respond to requested changes.", "Cannot review their own work or publish."],
            ["Reviewer", "Claim assigned or queued reviews, comment, approve, request changes or reject.", "Cannot edit the author’s draft or publish unless separately assigned Publisher."],
            ["Publisher", "See approved/scheduled work, publish now, schedule, reconcile and retract.", "Cannot alter the approved revision or record review decisions without Reviewer access."],
            ["Editorial Admin", "Supervise workload, roles, taxonomy, policies, audit, reviews and publishing operations.", "Immutable submitted, approved and published revisions remain protected."],
        ],
        [1900, 4550, 2910],
        font_size=9.0,
    )
    add_heading(doc, "What role-aware means", 2)
    paragraph(doc, "The same dashboard adapts to the signed-in person. Authors see their work, reviewers see eligible review tasks, publishers see content ready for delivery, and admins can supervise the whole system.", size=10.8)
    add_callout(doc, "Important", "The CEO/owner account is not intended as the everyday editorial admin profile. Normal client operations should use the dedicated roles above.")
    page_break(doc)

    add_page_title(doc, 3, "Creator structure and forms", "The workspace is the everyday interface; its forms are the auditable system of record.")
    add_table(
        doc,
        ["Area", "Core forms", "Purpose"],
        [
            ["Content", "Articles", "Master record: owner, category, policy, workflow state and active revision pointers."],
            ["Content", "Article Revisions", "Versioned content and SEO details; submitted, approved and published versions are protected."],
            ["Content", "Media Assets", "Cover and inline images with file details, alt text, caption, credit and processing status."],
            ["Discovery", "Categories and Tags", "Reusable taxonomy used in the editor, dashboard filters and public Insights archive."],
            ["People", "Employees and Role Assignments", "Maps users to active Author, Reviewer, Publisher and Admin capabilities."],
            ["Review", "Approval Policies", "Defines review model and number of distinct approvals required."],
            ["Review", "Review Assignments and Comments", "Reviewer slots, claims, decisions, summaries and discussion history."],
            ["Publishing", "Publication Jobs and Published Versions", "Tracks publish, schedule and retract actions plus the immutable Catalyst version."],
            ["Governance", "Audit Events, Redirects and Site Settings", "Preserves transitions and controls public behavior."],
        ],
        [1300, 2700, 5360],
        font_size=7.7,
    )
    add_callout(doc, "Client use", "Use forms for configuration and supervision; use the Article Workspace for lifecycle actions.")
    page_break(doc)

    add_page_title(doc, 4, "Dashboard orientation", "Start here to understand what needs attention, then narrow the workspace using scope, work type and filters.")
    add_figure(doc, dashboard, "Dashboard orientation — scope first, then work type, filters and task panels.", "Stylized dashboard showing Mine, Queue and All scopes, search, filters, summary metrics and work panels.")
    add_heading(doc, "The three scopes", 2)
    add_bullet(doc, "Mine — your authored articles and review assignments personally assigned to or claimed by you.")
    add_bullet(doc, "Queue — unclaimed shared reviews and publisher-ready work that your current role can act on.")
    add_bullet(doc, "All — cross-team visibility for Editorial Admin/owner supervision; unavailable to users without that authority.")
    add_heading(doc, "Find the right work quickly", 2)
    paragraph(doc, "Choose All work, Articles, Reviews or Publishing, then refine the result by state, category and sort order. Search can find titles, people, categories and review decisions. Clear filters returns to the default view without changing any content.", size=10.6)
    page_break(doc)

    add_page_title(doc, 5, "Create and prepare an article", "The writing workspace keeps editorial content, media, metadata and workflow context together.")
    add_heading(doc, "Start a draft", 2)
    add_numbered(doc, "Select New article from the dashboard.")
    add_numbered(doc, "Enter a working title, category and approval policy, then choose Create and open. Use Open full form only for administrative setup.")
    add_heading(doc, "Complete the writing workspace", 2)
    add_table(
        doc,
        ["Workspace area", "What to complete"],
        [
            ["Writing surface", "Body copy, H2–H4 headings, links, lists, quotes, tables and inline images."],
            ["Article settings", "Excerpt, category and tags."],
            ["Cover image", "JPG/PNG/WebP/GIF up to 10 MB, with required accessible alt text; caption and credit are optional."],
            ["Search settings", "SEO title, SEO description and robots directive."],
            ["Revision summary", "Current revision state, word count and estimated reading time."],
        ],
        [2300, 7060],
        font_size=8.8,
    )
    add_callout(doc, "Before submission", "Preview the article, confirm the cover and metadata, then submit the saved draft. The submitted revision becomes read-only so the reviewed content cannot silently change.")
    page_break(doc)

    add_page_title(doc, 6, "Choose the right review model", "Both models use independent reviewers and the same controlled publishing path.")
    add_figure(doc, review, "Standard and Regulated Review paths.", "Standard Review requires one distinct approval; Regulated Review requires two distinct reviewer approvals and remains In Review after the first.", width=4.4)
    add_heading(doc, "Reviewer assignment", 2)
    paragraph(doc, "The author can name eligible reviewers or use the shared queue; unfilled approval slots are queued automatically.", size=10.2)
    add_bullet(doc, "Self-review is blocked, and one reviewer cannot satisfy both Regulated Review slots.")
    add_bullet(doc, "Every claim, comment and decision is connected to the exact reviewed revision.")
    page_break(doc)

    add_page_title(doc, 7, "How a Standard Review article travels", "Use this path for normal editorial content that requires one independent approval.")
    add_numbered(doc, "Author drafts the article, completes media and metadata, and selects Standard Review.", num_id=52)
    add_numbered(doc, "Author submits to a named reviewer or the shared queue. The article becomes In Review and the submitted revision is locked.", num_id=52)
    add_numbered(doc, "Reviewer opens the Review Inbox, claims the assignment when required, reads the exact revision and adds discussion comments if needed.", num_id=52)
    add_numbered(doc, "Reviewer chooses Approve, Request Changes or Reject. Request Changes and Reject require a clear decision summary.", num_id=52)
    add_numbered(doc, "One approval moves the article to Approved. A Publisher can then publish immediately or schedule it.", num_id=52)
    add_heading(doc, "If changes are requested", 2)
    paragraph(doc, "The system preserves the submitted revision and creates Revision N+1 as a new Draft. The author sees the reviewer’s comments, updates the new draft and resubmits it through the same controlled path.", size=10.7)
    add_callout(doc, "Result", "The public site never receives a draft or an in-review revision—only the exact approved revision handed to publishing.")
    page_break(doc)

    add_page_title(doc, 8, "How a Regulated Review article travels", "Use this path when the content requires two independent approvals before it can reach publishing.")
    add_numbered(doc, "Author selects Regulated Review and submits the completed draft.", num_id=53)
    add_numbered(doc, "The system creates two reviewer slots using named reviewers, queue slots or a combination of both.", num_id=53)
    add_numbered(doc, "Reviewer A claims and approves the first slot. The article deliberately remains In Review.", num_id=53)
    add_numbered(doc, "Reviewer A cannot claim or approve the second slot for the same revision.", num_id=53)
    add_numbered(doc, "Reviewer B reviews independently and approves the second slot. The article becomes Approved exactly once.", num_id=53)
    add_numbered(doc, "Publisher completes immediate or scheduled publication using the same publishing boundary as Standard Review.", num_id=53)
    add_heading(doc, "Decision outcomes", 2)
    add_table(
        doc,
        ["Decision", "System behavior"],
        [
            ["Approve", "Records that reviewer’s independent decision; the article advances only when the policy threshold is met."],
            ["Request Changes", "Returns the article to the author, closes remaining open assignments and creates Draft Revision N+1."],
            ["Reject", "Stops the current review path with a required reason; admin policy determines any future reopening."],
        ],
        [2100, 7260],
        font_size=9.2,
    )
    page_break(doc)

    add_page_title(doc, 9, "Publishing and public delivery", "One approved revision becomes a durable public version through a secure, verified handoff.")
    add_figure(doc, publication, "Immediate and scheduled publication converge on the same verified handoff.", "Approved or scheduled content creates a Creator job, is validated and stored immutably by Catalyst, then appears on the public Insights site after a signed callback.", width=4.5)
    add_heading(doc, "Publisher controls", 2)
    add_bullet(doc, "Publish or Schedule — release the exact approved revision now or at a future time.")
    add_bullet(doc, "Reconcile — confirm a completed job if the final Creator callback has not settled.")
    add_bullet(doc, "Retract — stop public serving with a reason while preserving publication history.")
    add_callout(doc, "Safety", "Retries and reconciliation reuse the same job, preventing duplicate public versions.")
    page_break(doc)

    add_page_title(doc, 10, "Using the dashboard by role", "A practical starting view for each type of user.")
    add_table(
        doc,
        ["User", "Recommended starting view", "Typical next action"],
        [
            ["Author", "Mine → Articles", "Open a Draft, continue Changes Requested work or create a new article."],
            ["Reviewer", "Mine → Reviews, then Queue → Reviews", "Handle assigned work first; claim eligible shared-queue work when capacity allows."],
            ["Publisher", "Queue → Publishing", "Publish or schedule Approved items; monitor Processing, Scheduled and Needs retry states."],
            ["Editorial Admin", "All → All work", "Review bottlenecks, queue depth, changes requested, approved work and recent activity."],
        ],
        [1550, 3300, 4510],
        font_size=9.1,
    )
    add_heading(doc, "What the main panels mean", 2)
    add_bullet(doc, "Review Inbox — active assignments that are assigned, queued or claimed.")
    add_bullet(doc, "Article Work — drafts, changes requested and content currently in review.")
    add_bullet(doc, "Publishing Queue — approved, scheduled, processing or retryable publication work.")
    add_bullet(doc, "Published / Retracted — current public outcome with immutable history retained.")
    add_bullet(doc, "Review History and Activity — closed decisions and the chronological audit trail.")
    add_heading(doc, "Automatic email notifications", 2)
    paragraph(doc, "The platform sends best-effort, role-aware email updates after the related workflow action is safely recorded: reviewers are alerted when work is submitted; the author/owner is alerted when a review is claimed or decided; publishers are alerted when an article becomes ready; and the author/owner receives scheduling, publication, retraction or failure updates.", size=9.6, after=4)
    add_callout(doc, "Remember", "Email is a convenience alert. The dashboard, workflow state and Activity history remain the source of truth if an email is delayed or filtered.")
    add_callout(doc, "Live data", "The dashboard refreshes normally and checks more frequently while active publication jobs are running.")
    page_break(doc)

    add_page_title(doc, 11, "Status guide and good operating habits", "Use the state first: it tells you who owns the next step and which actions are safe.")
    add_table(
        doc,
        ["Status", "Meaning / next owner"],
        [
            ["Draft", "Editable working revision — Author."],
            ["In Review", "Submitted revision is locked; waiting for reviewer decisions — Reviewer(s)."],
            ["Changes Requested", "A new draft revision is available with feedback — Author."],
            ["Approved", "Review threshold is complete; waiting for publication — Publisher."],
            ["Scheduled / Processing", "Planned or active Catalyst delivery — Publisher monitors status and uses Reconcile only when prompted."],
            ["Published", "Visible on the public Insights site — Publisher/Admin monitors."],
            ["Unpublished / Retracted", "Removed from public serving; immutable history remains — Publisher/Admin."],
        ],
        [2300, 7060],
        font_size=7.8,
    )
    add_heading(doc, "Recommended habits", 2)
    add_bullet(doc, "Use the Article Workspace for lifecycle actions; avoid editing workflow records directly.")
    add_bullet(doc, "Before submission, confirm the policy, reviewers, alt text and SEO metadata.")
    add_bullet(doc, "Read the decision summary before responding; reconcile only when prompted and retract rather than delete.")
    add_callout(doc, "When unsure", "Confirm the article’s state and revision, then read Workflow facts and Activity for the next step.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
